from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "Trigger_Box_Controller_Explanation_v000.docx"
FLOWCHART_PATH = ROOT / "trigger_box_sequence_diagram_v002.png"
FALLBACK_FLOWCHART_PATH = ROOT / "trigger_box_sequence_diagram.png"


ACCENT = RGBColor(27, 76, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "EEF4F8"
LIGHT_BORDER = "D9E2EA"


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)


def add_number(document, text):
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)


def create_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = ACCENT

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(11.5)
    subtitle.font.color.rgb = MUTED

    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT

    if "Callout" not in styles:
        style = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_p.add_run("Trigger Box Controller Guide")
    header_run.font.name = "Arial"
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = MUTED

    p_pr = header_p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFCAD4")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_p.add_run("Page ")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = MUTED
    set_page_number(footer_p)


def add_paragraph(document, text, bold_prefix=None):
    p = document.add_paragraph()
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def enforce_run_font(run, font_name="Arial"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def enforce_arial_everywhere(document):
    def process_paragraphs(paragraphs):
        for paragraph in paragraphs:
            for run in paragraph.runs:
                enforce_run_font(run)

    process_paragraphs(document.paragraphs)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    for section in document.sections:
        process_paragraphs(section.header.paragraphs)
        process_paragraphs(section.footer.paragraphs)


def build_flowchart_image(path: Path):
    width, height = 1800, 1200
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    try:
        title_font = ImageFont.truetype("Arial.ttf", 40)
        section_font = ImageFont.truetype("Arial Bold.ttf", 24)
        box_font = ImageFont.truetype("Arial.ttf", 24)
        small_font = ImageFont.truetype("Arial.ttf", 20)
    except OSError:
        title_font = ImageFont.load_default()
        section_font = ImageFont.load_default()
        box_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    def box(x1, y1, x2, y2, text, fill="#EEF4F8", outline="#1B4C78", font=box_font, radius=22):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=radius, fill=fill, outline=outline, width=3)
        draw.multiline_text(
            ((x1 + x2) / 2, (y1 + y2) / 2),
            text,
            fill="#111111",
            font=font,
            anchor="mm",
            align="center",
            spacing=4,
        )

    def arrow(x1, y1, x2, y2, text=None, color="#48515A", width_px=4):
        draw.line((x1, y1, x2, y2), fill=color, width=width_px)
        if abs(y2 - y1) >= abs(x2 - x1):
            direction = 1 if y2 > y1 else -1
            tip = [(x2, y2), (x2 - 10, y2 - 18 * direction), (x2 + 10, y2 - 18 * direction)]
        else:
            direction = 1 if x2 > x1 else -1
            tip = [(x2, y2), (x2 - 18 * direction, y2 - 10), (x2 - 18 * direction, y2 + 10)]
        draw.polygon(tip, fill=color)
        if text:
            draw.rounded_rectangle(
                (((x1 + x2) / 2) - 90, ((y1 + y2) / 2) - 22, ((x1 + x2) / 2) + 90, ((y1 + y2) / 2) + 10),
                radius=10,
                fill="white",
            )
            draw.text(((x1 + x2) / 2, (y1 + y2) / 2 - 6), text, fill=color, font=small_font, anchor="mm")

    def connector(points, text=None, color="#48515A"):
        for start, end in zip(points, points[1:]):
            arrow(start[0], start[1], end[0], end[1], None, color=color)
        if text:
            mid = points[len(points) // 2]
            draw.rounded_rectangle((mid[0] - 95, mid[1] - 24, mid[0] + 95, mid[1] + 8), radius=10, fill="white")
            draw.text((mid[0], mid[1] - 8), text, fill=color, font=small_font, anchor="mm")

    def bullet_list(x, y, lines, color="#2A2A2A"):
        cursor_y = y
        for line in lines:
            draw.ellipse((x, cursor_y + 8, x + 8, cursor_y + 16), fill=color)
            draw.text((x + 24, cursor_y), line, fill=color, font=small_font)
            cursor_y += 36

    draw.rounded_rectangle((40, 40, width - 40, height - 40), radius=30, outline="#D8E0E8", width=2)
    draw.rounded_rectangle((65, 95, width - 65, 150), radius=20, fill="#F5F8FB", outline="#D8E0E8", width=2)
    draw.text((width / 2, 122), "Trigger Box Sequence Overview", fill="#1B4C78", font=title_font, anchor="mm")

    draw.text((330, 200), "Spark-test branch", fill="#1B4C78", font=section_font, anchor="mm")
    draw.text((1110, 200), "Hydrogen-test branch", fill="#1B4C78", font=section_font, anchor="mm")
    draw.text((1560, 200), "Lockout / reset", fill="#9A5B00", font=section_font, anchor="mm")

    box(720, 230, 1080, 310, "Power on\nSafe startup", fill="#F7FAFC")
    box(720, 360, 1080, 440, "Mode selection", fill="#F7FAFC")
    arrow(900, 310, 900, 360)

    box(90, 500, 560, 670, "Spark-test mode\n\nArm active -> periodic spark pulses\nHot wires OFF\nDAQ OFF")
    connector([(720, 400), (560, 400), (325, 500)], "Mode HIGH", color="#1B4C78")

    box(720, 500, 1080, 580, "Idle\nAll outputs OFF")
    box(720, 640, 1080, 720, "Armed\nArm light ON\nWaiting for Trigger")
    box(720, 780, 1080, 880, "Melting\n3 hot-wire relays ON\nfor hotWireBurn_us")
    box(720, 940, 1080, 1020, "Spark ON\nDAQ still OFF")
    box(1180, 940, 1540, 1020, "DAQ ON during final\ndaqPulse_us of spark", fill="#EEF7F2", outline="#2E7D4F")
    box(1180, 1080, 1540, 1160, "Fired lockout\nAll outputs OFF", fill="#F3F6FA")

    connector([(1080, 400), (1260, 400), (1260, 470), (900, 470), (900, 500)], "Mode LOW", color="#1B4C78")
    arrow(900, 580, 900, 640, "Arm active")
    arrow(900, 720, 900, 780, "Trigger")
    arrow(900, 880, 900, 940)
    arrow(1080, 980, 1180, 980, "start DAQ", color="#2E7D4F")
    arrow(1360, 1020, 1360, 1080, "end spark", color="#2E7D4F")

    box(1380, 500, 1720, 690, "Lockout states\n\nFAIL\nABORTED\nFIRED", fill="#FFF4E8", outline="#9A5B00")
    box(1380, 760, 1720, 920, "Reset condition\n\nArm = OFF\nTrigger = OFF", fill="#FFF9F0", outline="#9A5B00")

    img.save(path)


def build_document():
    document = Document()
    create_styles(document)
    configure_page(document)

    if not FLOWCHART_PATH.exists():
        build_flowchart_image(FALLBACK_FLOWCHART_PATH)

    title = document.add_paragraph(style="Title")
    title.add_run("Trigger Box Controller")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Code review summary of the reviewed M-Duino control logic for operator understanding, peer review, and lab handover."
    )

    meta = document.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = True
    entries = [
        ("Target platform", "Industrial Shields M-Duino 19R+"),
        ("Document date", str(date.today())),
        ("Purpose", "Explain the code structure, firing sequence, timing variables, and safety behavior in words."),
    ]
    for r, (label, value) in enumerate(entries):
        set_cell_text(meta.cell(r, 0), label, bold=True)
        set_cell_text(meta.cell(r, 1), value)
    style_table(meta)

    document.add_heading("1. What This Code Does", level=1)
    add_paragraph(
        document,
        "The controller manages a test box with two operating modes. In spark-test mode it repeatedly produces short spark pulses while the Arm input remains active. In hydrogen-test mode it enforces a stricter sequence: the operator arms the system, presses Trigger once, the hot-wire relays energize for a defined time, the spark output fires, and the DAQ trigger turns on near the end of the spark pulse.",
    )
    add_paragraph(
        document,
        "The revised code is longer than the original because it makes the sequence explicit. Instead of relying on several interacting flags, it uses a named state machine so that anyone reading the code can tell what the controller is doing at any moment.",
    )

    document.add_heading("2. Main Hardware Signals", level=1)
    io_table = document.add_table(rows=1, cols=4)
    headers = ["Signal", "Type", "Used for", "Meaning in the logic"]
    for i, header in enumerate(headers):
        set_cell_text(io_table.cell(0, i), header, bold=True)
    rows = [
        ("Arm input", "Digital input", "Both modes", "Allows the system to enter an armed state and must remain active during hazardous phases."),
        ("Trigger input", "Digital input", "Hydrogen-test mode", "Starts the firing sequence after arming. It is treated as a momentary start command."),
        ("Mode input", "Digital input", "Both modes", "Selects spark-test mode or hydrogen-test mode."),
        ("Arm light", "Digital output", "Both modes", "Shows that the controller is armed or in an active sequence."),
        ("HotWire 1/2/3", "Digital outputs", "Hydrogen-test mode", "Drive the three relay control lines for the external hot-wire power path."),
        ("SparkOut", "Digital output", "Both modes", "Commands the spark or ignition stage."),
        ("DAQTrig", "Digital output", "Hydrogen-test mode", "Provides a timing signal to the data-acquisition system."),
    ]
    for row in rows:
        cells = io_table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, bold=False)
    style_table(io_table)

    document.add_heading("3. Important Unit Convention", level=1)
    add_paragraph(
        document,
        "The revised code keeps units visible in the variable names wherever units matter. Any variable ending in _us is a time in microseconds. This makes the timing easier to audit and reduces confusion when the values are small.",
    )
    add_bullet(document, "hotWireBurn_us: hot-wire relay on-time before spark begins")
    add_bullet(document, "sparkDwell_us: total spark pulse duration")
    add_bullet(document, "daqPulse_us: DAQ pulse width during the end of the spark pulse")
    add_bullet(document, "sparkTestInterval_us: time between spark pulses in spark-test mode")
    add_bullet(document, "debounce_us: time used to reject switch bounce")
    add_bullet(document, "lastSerialPrint_us, meltStart_us, sparkStart_us: timestamps stored with micros()")
    add_paragraph(
        document,
        "Pins, booleans, and states do not carry physical units, so they intentionally do not use a unit suffix.",
    )

    document.add_heading("4. Sequence of Events", level=1)
    add_paragraph(
        document,
        "The figure below summarizes the expected sequence. The left branch shows spark-test mode. The right branch shows the more controlled hydrogen-test mode with lockout behavior.",
    )
    picture_path = FLOWCHART_PATH if FLOWCHART_PATH.exists() else FALLBACK_FLOWCHART_PATH
    document.add_picture(str(picture_path), width=Inches(6.6))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run("Figure 1. Overview of the trigger box sequence.")
    cap_run.italic = True
    cap_run.font.size = Pt(9.5)
    cap_run.font.color.rgb = MUTED

    document.add_heading("5. Hydrogen-Test Sequence in Words", level=1)
    add_number(document, "The controller starts in a safe idle state with all outputs turned off.")
    add_number(document, "If Trigger is already active before proper arming, the controller enters a fail lockout state. This forces the operator to release both Arm and Trigger before trying again.")
    add_number(document, "When Arm becomes active, the controller turns on the Arm light and enters the armed state.")
    add_number(document, "When Trigger is pressed, the controller either starts the hot-wire stage or jumps directly to spark if the hot-wire step has been disabled in the code.")
    add_number(document, "During the melting stage, all three hot-wire relay outputs are energized for hotWireBurn_us microseconds.")
    add_number(document, "After the hot-wire time expires, the controller turns the hot wires off and starts the spark pulse.")
    add_number(document, "The DAQ output remains off at first, then turns on during only the last daqPulse_us microseconds of the spark pulse.")
    add_number(document, "At the end of the spark dwell time, both SparkOut and DAQTrig are turned off and the controller enters a fired lockout state.")
    add_number(document, "The controller stays locked out until both Arm and Trigger have been released. Only then does it return to idle.")

    document.add_heading("6. Spark-Test Mode in Words", level=1)
    add_paragraph(
        document,
        "Spark-test mode is intentionally simpler. The hot-wire outputs remain off and the DAQ output remains off. If Arm is active, the controller produces periodic spark pulses. If Arm is released, the Arm light and SparkOut both turn off immediately.",
    )
    add_paragraph(
        document,
        "This mode is useful for checking the ignition side without energizing the hot-wire relays or involving the DAQ timing logic.",
    )

    document.add_heading("7. Why the State Machine Matters", level=1)
    add_paragraph(
        document,
        "The original short sketch used several flags such as Armed, Fired, Fail, and flag. That works for quick experiments, but it becomes hard to see what the controller should do after unusual operator actions. The reviewed version uses named states instead: Idle, Armed, Melting, SparkWaitDaq, SparkWaitEnd, Fired, Fail, and Aborted.",
    )
    add_paragraph(
        document,
        "This makes debugging easier because the code always answers the question: what state is the controller in right now? It also makes the serial output more useful, since the current state can be printed directly as text.",
    )

    document.add_heading("8. Timing Variables and Their Meaning", level=1)
    timing = document.add_table(rows=1, cols=4)
    for idx, header in enumerate(["Variable", "Current value", "Unit", "Practical meaning"]):
        set_cell_text(timing.cell(0, idx), header, bold=True)
    timing_rows = [
        ("hotWireBurn_us", "10,000,000", "microseconds", "Keeps the three hot-wire relays energized for 10 seconds."),
        ("sparkDwell_us", "5,000", "microseconds", "Length of the spark command pulse."),
        ("daqPulse_us", "600", "microseconds", "Width of the DAQ trigger pulse at the end of the spark pulse."),
        ("sparkTestInterval_us", "500,000", "microseconds", "Time between spark pulses in spark-test mode."),
        ("debounce_us", "30,000", "microseconds", "Minimum stable input time before a switch change is accepted."),
        ("serialPrintInterval_us", "250,000", "microseconds", "Limits how often status lines are sent to the serial monitor."),
    ]
    for row in timing_rows:
        cells = timing.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, bold=False)
    style_table(timing)

    document.add_heading("9. Safety Behaviors Built Into the Logic", level=1)
    add_bullet(document, "All outputs are forced off at startup.")
    add_bullet(document, "Changing Mode during operation forces a safe reset back to idle.")
    add_bullet(document, "If Arm is released during melting or spark, the controller aborts immediately and enters lockout.")
    add_bullet(document, "After a fired, failed, or aborted sequence, the operator must release both Arm and Trigger before restarting.")
    add_bullet(document, "The hot-wire step and spark step are never left on together accidentally; hot wires are forced off before spark begins.")

    document.add_heading("10. Code Structure", level=1)
    structure = document.add_table(rows=1, cols=2)
    for idx, header in enumerate(["Function or block", "Role in the program"]):
        set_cell_text(structure.cell(0, idx), header, bold=True)
    structure_rows = [
        ("setup()", "Configures inputs and outputs, starts serial communication, and forces a safe startup state."),
        ("loop()", "Updates inputs, checks for mode changes, runs the appropriate mode handler, and prints status."),
        ("DebouncedInput", "Filters switch bounce so mechanical inputs behave more reliably."),
        ("handleSparkTestMode()", "Runs the periodic spark-only behavior."),
        ("handleHydrogenTestMode()", "Runs the main state machine for arming, melting, spark, DAQ timing, and lockout."),
        ("startMeltingOrSpark()", "Chooses whether to start the hot-wire stage or skip directly to spark."),
        ("startSparkSequence()", "Turns hot wires off, enables spark, and starts spark timing."),
        ("allOutputsOff()", "Provides one reusable safe-off command for all controlled outputs."),
    ]
    for row in structure_rows:
        cells = structure.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    style_table(structure)

    document.add_heading("11. Two Hardware Points To Confirm", level=1)
    add_bullet(document, "Relay polarity: if the relay modules energize when the M-Duino output goes LOW, the output logic must be inverted in the code.")
    add_bullet(document, "Input wiring: if the switches are not electrically well-defined, the input logic may float. The wiring and pull-up or pull-down strategy should match the real M-Duino hardware.")

    document.add_heading("12. Suggested Lab Use", level=1)
    add_number(document, "Test the controller first with the Mac connected by USB and with the hazardous power stages disconnected.")
    add_number(document, "Watch the serial monitor to confirm Mode, Arm, Trigger, and state transitions.")
    add_number(document, "Verify relay polarity on each hot-wire channel before connecting the real hot-wire power source.")
    add_number(document, "Only after the low-risk checks pass should the real hot-wire and spark hardware be introduced.")

    document.add_heading("13. Short Summary", level=1)
    add_paragraph(
        document,
        "In plain terms, the reviewed code is a safer and easier-to-read version of the original trigger sketch. It keeps the same main purpose, but it explains the timing clearly, separates relay outputs from boolean options, and makes the firing sequence explicit enough to review with colleagues.",
    )

    enforce_arial_everywhere(document)
    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
