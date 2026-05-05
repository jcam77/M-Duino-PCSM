from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent.parent
DOCX_PATH = ROOT / "M-Duino-Trigger_Box_Code_Explanation" / "Original_Trigger_Box_Code_Explanation_v000.docx"
FLOWCHART_PATH = ROOT / "Diagrams" / "original_trigger_box_expected_sequence_v000.png"

ACCENT = RGBColor(27, 76, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "EEF4F8"


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def enforce_run_font(run, font_name="Arial"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def enforce_arial_everywhere(document):
    def process_paragraphs(paragraphs):
        for paragraph in paragraphs:
            for run in paragraph.runs:
                enforce_run_font(run)

    process_paragraphs(document.paragraphs)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    for section in document.sections:
        process_paragraphs(section.header.paragraphs)
        process_paragraphs(section.footer.paragraphs)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    enforce_run_font(run)
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def add_number(document, text):
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def create_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = ACCENT

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(11.5)
    subtitle.font.color.rgb = MUTED

    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT

    if "Callout" not in styles:
        style = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["Normal"]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_p.add_run("Original Trigger Box Code Review")
    enforce_run_font(header_run)
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = MUTED

    p_pr = header_p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFCAD4")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_p.add_run("Page ")
    enforce_run_font(footer_run)
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = MUTED
    set_page_number(footer_p)


def add_paragraph(document, text, bold_prefix=None):
    p = document.add_paragraph()
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        enforce_run_font(run)
        run2 = p.add_run(text)
        enforce_run_font(run2)
    else:
        run = p.add_run(text)
        enforce_run_font(run)
    return p


def build_flowchart_image(path: Path):
    width, height = 1800, 1200
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    try:
        title_font = ImageFont.truetype("Arial Bold.ttf", 44)
        section_font = ImageFont.truetype("Arial Bold.ttf", 24)
        box_font = ImageFont.truetype("Arial.ttf", 24)
        small_font = ImageFont.truetype("Arial.ttf", 20)
    except OSError:
        title_font = ImageFont.load_default()
        section_font = ImageFont.load_default()
        box_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    def box(x1, y1, x2, y2, text, fill="#F7FAFC", outline="#1B4C78"):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=fill, outline=outline, width=3)
        draw.multiline_text(((x1 + x2) / 2, (y1 + y2) / 2), text, font=box_font, fill="#111111", anchor="mm", align="center", spacing=4)

    def arrow(x1, y1, x2, y2, text=None, color="#274E7A"):
        draw.line((x1, y1, x2, y2), fill=color, width=4)
        if abs(y2 - y1) >= abs(x2 - x1):
            direction = 1 if y2 > y1 else -1
            tip = [(x2, y2), (x2 - 11, y2 - 18 * direction), (x2 + 11, y2 - 18 * direction)]
        else:
            direction = 1 if x2 > x1 else -1
            tip = [(x2, y2), (x2 - 18 * direction, y2 - 11), (x2 - 18 * direction, y2 + 11)]
        draw.polygon(tip, fill=color)
        if text:
            draw.rounded_rectangle((((x1 + x2) / 2) - 85, ((y1 + y2) / 2) - 20, ((x1 + x2) / 2) + 85, ((y1 + y2) / 2) + 8), radius=10, fill="white")
            draw.text(((x1 + x2) / 2, (y1 + y2) / 2 - 6), text, font=small_font, fill=color, anchor="mm")

    draw.rounded_rectangle((40, 40, width - 40, height - 40), radius=28, outline="#D8E0E8", width=2)
    draw.rounded_rectangle((65, 95, width - 65, 150), radius=18, fill="#F5F8FB", outline="#D8E0E8", width=2)
    draw.text((width / 2, 122), "Original Code: Expected Sequence", font=title_font, fill="#1B4C78", anchor="mm")

    draw.text((360, 205), "Spark-test branch", font=section_font, fill="#1B4C78", anchor="mm")
    draw.text((1020, 205), "Hydrogen-test branch", font=section_font, fill="#1B4C78", anchor="mm")
    draw.text((1510, 205), "Observed risks", font=section_font, fill="#9A5B00", anchor="mm")

    box(730, 240, 1090, 320, "Mode check")
    box(120, 450, 600, 610, "ConstSpark()\n\nIf Arm = HIGH:\nSparkOut pulses for dwell\nArmLight ON")
    arrow(730, 280, 600, 520, "Mode = HIGH")

    box(730, 400, 1090, 480, "Idle / reset\nTrigger = LOW\nArm = LOW")
    box(730, 560, 1090, 640, "Armed\nArm = HIGH\nFail = false")
    box(730, 720, 1090, 820, "Melty()\n3 hot wires ON\nfor BurnTime x delay(1000)")
    box(730, 900, 1090, 980, "SparkOut HIGH\nthen DAQTrig HIGH\nthen both LOW")
    box(730, 1060, 1090, 1140, "Fired latch\nmanual reset required")

    arrow(910, 480, 910, 560, "Arm = HIGH")
    arrow(910, 640, 910, 720, "Trigger = HIGH")
    arrow(910, 820, 910, 900, "after Melty()")
    arrow(910, 980, 910, 1060, "Fired = true")

    box(1310, 360, 1710, 520, "Issue 1\nNo comments in the original file\nso units and intent are hidden", fill="#FFF4E8", outline="#9A5B00")
    box(1310, 560, 1710, 720, "Issue 2\nHotWire is a boolean flag\nbut is used like a pin", fill="#FFF4E8", outline="#9A5B00")
    box(1310, 760, 1710, 920, "Issue 3\nHotWire1/2/3 are used\nbut never set with pinMode(..., OUTPUT)", fill="#FFF4E8", outline="#9A5B00")
    box(1310, 960, 1710, 1120, "Issue 4\nBlocking delays mean the controller\ncannot react during Melty() or spark timing", fill="#FFF4E8", outline="#9A5B00")

    img.save(path)


def build_document():
    document = Document()
    create_styles(document)
    configure_page(document)
    build_flowchart_image(FLOWCHART_PATH)

    title = document.add_paragraph(style="Title")
    run = title.add_run("Original Trigger Box Code")
    enforce_run_font(run)

    subtitle = document.add_paragraph(style="Subtitle")
    run = subtitle.add_run(
        "Code review summary of the original M-Duino sketch, including intended behavior, missing documentation, and implementation risks."
    )
    enforce_run_font(run)

    meta = document.add_table(rows=4, cols=2)
    entries = [
        ("Code file", "M-DuinoScripts/M-Duino_Original/M-Duino_Original.ino"),
        ("Document date", str(date.today())),
        ("Focus", "Explain what the original code appears to do and identify the main technical risks."),
        ("Important note", "This document describes the original sketch, not the later cleaned state-machine review path that currently reaches M_Duino_v005.ino."),
    ]
    for r, (label, value) in enumerate(entries):
        set_cell_text(meta.cell(r, 0), label, bold=True)
        set_cell_text(meta.cell(r, 1), value)
    style_table(meta)

    document.add_heading("1. Why This Document Exists", level=1)
    add_paragraph(document, "The original trigger-box sketch is short, but that brevity hides several important assumptions. The file contains almost no comments, several variable names do not show their units, and at least two lines treat a boolean flag as if it were a hardware output pin. This document explains the original code in words so that colleagues can understand what it seems intended to do and why it is risky to use without review.")

    document.add_heading("2. Code File Covered Here", level=1)
    add_bullet(document, "M-DuinoScripts/M-Duino_Original/M-Duino_Original.ino")
    add_paragraph(document, "Important note:")
    add_bullet(document, "This document describes the original sketch only.")
    add_bullet(document, "It does not describe the later cleaned state-machine review path that currently reaches `M_Duino_v005.ino`.")
    add_bullet(document, "The original `.ino` file is the source of truth for the baseline implementation discussed here.")

    document.add_heading("3. Relationship to the Later Reviewed Version", level=1)
    add_paragraph(document, "This document describes the original baseline code that was later improved in the reviewed firmware.")
    add_paragraph(document, "The current reviewed production version is:")
    add_bullet(document, "M-DuinoScripts/M_Duino_v005/M_Duino_v005.ino")
    add_paragraph(document, "So the intended interpretation is:")
    add_bullet(document, "this document explains the original implementation and its risks")
    add_bullet(document, "the reviewed controller document explains the later improved version")
    add_bullet(document, "the two documents should not be treated as describing the same firmware file")

    document.add_heading("4. What the Original Code Appears to Do", level=1)
    add_paragraph(document, "The original sketch has two modes. When Mode is HIGH, it runs ConstSpark(), which repeatedly produces a spark pulse whenever Arm is active. When Mode is LOW, it runs HydrogenTest(), which appears intended to enforce an arming sequence, activate the hot-wire relays, fire the spark, issue a DAQ pulse, and then latch the system as fired until the operator resets it.")
    add_paragraph(document, "Although this intended sequence is visible by reading the code carefully, it is not written down anywhere in the original file. The next figure summarizes the expected behavior the code seems to be aiming for.")

    document.add_heading("5. Expected Sequence from the Original Logic", level=1)
    document.add_picture(str(FLOWCHART_PATH), width=Inches(6.6))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run("Figure 1. Expected sequence and major risks in the original sketch.")
    enforce_run_font(run)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED

    document.add_heading("6. Missing Comments and Hidden Assumptions", level=1)
    add_bullet(document, "There are no top-level comments explaining the operating modes, safety assumptions, or intended firing sequence.")
    add_bullet(document, "Variables such as BurnTime, dwell, and Delay do not carry units in their names, so the reader must infer the units from delay() and delayMicroseconds().")
    add_bullet(document, "The code does not explain whether the inputs and outputs are active HIGH or active LOW.")
    add_bullet(document, "The code does not explain that the hot-wire outputs are assumed to drive relays or another external switching stage rather than the hot wire directly.")
    add_paragraph(document, "Because of these missing comments, different readers can misunderstand the same code in different ways. That is the main reason the cleaned rewrite adds comments, unit suffixes, and explicit state names.")

    document.add_heading("7. Hidden Time Units in the Original Code", level=1)
    add_paragraph(document, "The original code works with mixed time scales, but the units are hidden. The M-Duino does not know the time unit from the number alone. It knows the time unit from the function the number is passed into.")

    table = document.add_table(rows=1, cols=4)
    for idx, header in enumerate(["Original variable", "Current value", "Effective unit", "How the unit is inferred"]):
        set_cell_text(table.cell(0, idx), header, bold=True)
    rows = [
        ("BurnTime", "10", "seconds in practice", "Used as the repeat count for delay(1000), so 10 x 1000 ms = 10 s."),
        ("dwell", "5000", "microseconds", "Used in delayMicroseconds(dwell)."),
        ("Delay", "600", "microseconds", "Used in delayMicroseconds(Delay)."),
        ("count", "loop counter", "no physical unit", "Used only to count the BurnTime loop."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    style_table(table)

    document.add_heading("8. Main Technical Issues in the Original Sketch", level=1)
    issues = document.add_table(rows=1, cols=3)
    for idx, header in enumerate(["Issue", "Where it appears", "Why it matters"]):
        set_cell_text(issues.cell(0, idx), header, bold=True)
    issue_rows = [
        ("Boolean used as pin", "pinMode(HotWire, OUTPUT) and digitalWrite(HotWire, LOW)", "HotWire is declared as a boolean flag, so this code is not controlling HotWire1/2/3 as intended."),
        ("Hot-wire outputs not configured", "Melty() uses HotWire1/2/3 but setup() never calls pinMode() for them", "Relay behavior may be unreliable because the outputs are not explicitly configured."),
        ("Blocking delays", "Melty(), ConstSpark(), and the spark/DAQ timing block", "The controller cannot check Arm or any abort condition while it is stuck inside these delays."),
        ("No explicit state machine", "HydrogenTest() mixes Fail, Fired, Armed, and flag", "It is hard to tell what the controller should do after unusual operator actions."),
        ("Repeated raw digitalRead calls", "HydrogenTest() and loop()", "The logic is harder to follow and there is no debounce handling."),
        ("Very limited serial output", "loop() only prints raw input/state bits", "The operator cannot directly see state transitions such as arming, melting, firing, or fail conditions."),
    ]
    for row in issue_rows:
        cells = issues.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    style_table(issues)

    document.add_heading("9. Original Hydrogen-Test Logic in Words", level=1)
    add_number(document, "If both Trigger and Arm are LOW, the code clears Fired and Fail.")
    add_number(document, "If Trigger is HIGH while Arm is LOW, the code sets Fail = true.")
    add_number(document, "If Fail is false, Arm is HIGH, and Fired is false, the code sets Armed = true and turns on the ArmLight.")
    add_number(document, "If Trigger is then HIGH while Armed is true and Fired is false, the code enters the firing block.")
    add_number(document, "Inside the firing block, if HotWire is true, the code calls Melty(), which turns on the three hot-wire outputs for BurnTime cycles of delay(1000).")
    add_number(document, "After Melty(), the code turns SparkOut on, waits dwell - Delay microseconds, turns DAQTrig on, waits Delay microseconds, then turns both outputs off.")
    add_number(document, "The code sets Fired = true and then uses flag as an additional latch.")
    add_number(document, "After firing, the operator must release both Arm and Trigger to clear the Fired and Fail conditions.")

    document.add_heading("10. Why the Original File Is Hard to Debug", level=1)
    add_paragraph(document, "The original file does not have named states such as IDLE, ARMED, MELTING, or FIRED. Instead, the operator has to infer the current situation from several booleans that interact with each other. This makes it hard to answer questions like: Did the sequence fail before arming? Is it currently in the hot-wire phase? Has it already fired once? Is it waiting for reset?")
    add_paragraph(document, "In addition, the serial output in the original sketch prints compact numeric values without describing transitions in words. That means the code may be running, but the operator still cannot easily see the sequence in real time.")

    document.add_heading("11. Safety Concerns Specific to the Original Code", level=1)
    add_bullet(document, "There is no hardwired safety logic in software; the file assumes external protection exists.")
    add_bullet(document, "Because Melty() blocks for the full hot-wire duration, the code cannot react to Arm being released during that period.")
    add_bullet(document, "The original file contains no clear documentation of safe reset conditions or mode-change behavior.")
    add_bullet(document, "The hot-wire relay control is mixed up with a boolean flag, which is especially dangerous in a hardware-control file because it can create false confidence about what output is being driven.")

    document.add_heading("12. Practical Reading of the Original File", level=1)
    add_paragraph(document, "The original sketch should be read as a first working prototype rather than a finished control program. It captures the intended test sequence in a compact form, but it leaves too much hidden in the implementation. For collaboration, safety review, and commissioning, the code benefits from explicit comments, visible units, deliberate debug output, and a clearer state structure.")

    document.add_heading("13. Short Summary", level=1)
    add_paragraph(document, "The original code appears intended to run a valid trigger-box sequence, but it contains missing comments, hidden units, blocking timing, and at least one major pin/boolean mix-up. The code can therefore be understood as a prototype of the desired sequence, but not as a well-documented or low-risk final implementation.")

    enforce_arial_everywhere(document)
    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
