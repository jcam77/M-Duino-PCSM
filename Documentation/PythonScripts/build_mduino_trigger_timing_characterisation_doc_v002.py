from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
DOCX_PATH = ROOT / "M-Duino_Trigger_Timing_Characterisation_Guide" / "M-Duino_Trigger_Timing_Characterisation_Guide_v002.docx"
DIAGRAM_PATH = ROOT / "Diagrams" / "mduino_trigger_scope_connections_v000.png"

ACCENT = RGBColor(27, 76, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "EEF4F8"


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def enforce_run_font(run, font_name="Arial"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    enforce_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)


def create_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = ACCENT

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(11.5)
    subtitle.font.color.rgb = MUTED

    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_p.add_run("M-Duino Trigger Timing Characterisation Guide")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    p_pr = header_p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFCAD4")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("Page ")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    set_page_number(footer_p)


def add_paragraph(document, text):
    p = document.add_paragraph()
    run = p.add_run(text)
    enforce_run_font(run)
    return p


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def add_number(document, text):
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def build_document():
    document = Document()
    create_styles(document)
    configure_page(document)

    title = document.add_paragraph(style="Title")
    run = title.add_run("M-Duino Trigger Timing Characterisation Guide")
    enforce_run_font(run)

    subtitle = document.add_paragraph(style="Subtitle")
    run = subtitle.add_run("Detailed oscilloscope procedure with connection diagrams for delay, jitter, and synchronisation checks.")
    enforce_run_font(run)

    meta = document.add_table(rows=4, cols=2)
    entries = [
        ("Document date", str(date.today())),
        ("Application", "M-Duino trigger coordination"),
        ("Primary use", "Deflagration trigger-path characterisation"),
        ("Focus", "Delay, jitter, relative skew, and connection setup"),
    ]
    for r, (label, value) in enumerate(entries):
        set_cell_text(meta.cell(r, 0), label, bold=True)
        set_cell_text(meta.cell(r, 1), value)
    style_table(meta)

    document.add_heading("1. Purpose", level=1)
    add_paragraph(document, "This guide explains how to measure the timing behaviour of the M-Duino when it is used as a trigger coordinator for a hydrogen deflagration setup. The objective is to quantify controller delay, shot-to-shot jitter, and the relative synchronisation offsets between the DAQ trigger branch and the camera trigger branch.")

    document.add_heading("2. Connection Diagram", level=1)
    add_paragraph(document, "The figure below shows the recommended oscilloscope hookup logic for the three main measurement runs.")
    if DIAGRAM_PATH.exists():
        document.add_picture(str(DIAGRAM_PATH), width=Inches(7.35))
    add_paragraph(document, "Use each panel as a separate measurement configuration. Do not try to measure everything at once if doing so makes the probing unclear or unreliable.")

    document.add_heading("3. What This Characterisation Is Trying To Prove", level=1)
    add_paragraph(document, "The M-Duino is not the final scientific timing reference in this setup. The final scientific time base is provided by the high-speed datalogger. The main question is therefore whether the M-Duino distributes the trigger events with low enough uncertainty for the experiment.")
    for item in [
        "Low and stable trigger delay",
        "Low trigger jitter",
        "Low relative skew between systems",
    ]:
        add_bullet(document, item)

    document.add_heading("4. Assumed Trigger Architecture", level=1)
    for item in [
        "A clean input event is applied to the M-Duino.",
        "The M-Duino produces a DAQ trigger output.",
        "That DAQ trigger is split through a BNC splitter to DAQ 1 and DAQ 2.",
        "The camera is triggered through its own separate trigger path.",
        "Final scientific timestamping is performed by the high-speed datalogger.",
    ]:
        add_number(document, item)

    document.add_heading("5. Why Each Measurement Is Needed", level=1)
    for item in [
        "Input to M-Duino DAQ trigger output: controller reaction delay and jitter.",
        "Input to camera trigger path: controller-to-camera path timing and jitter.",
        "M-Duino DAQ output to DAQ 1 and DAQ 2 inputs: splitter-path delay and branch-to-branch comparison.",
        "Camera trigger path relative to the DAQ trigger path: inter-system synchronisation offset and stability.",
    ]:
        add_bullet(document, item)

    document.add_heading("6. Do You Need Special Benchmark Code?", level=1)
    add_paragraph(document, "Not necessarily for the first pass. A real-system timing study with the actual firmware and actual trigger paths is usually the most important measurement because it answers whether the experiment-ready system is good enough.")
    add_paragraph(document, "Special benchmark firmware of the type shown in the Industrial Shields article is optional and mainly useful if the real-system timing looks worse than expected or if the delay needs to be decomposed into hardware and firmware contributions.")

    document.add_heading("7. Safety and Measurement Validity", level=1)
    for item in [
        "Perform the test without ignition and without hazardous experiment execution.",
        "Use a safe and repeatable pulse source.",
        "Use the same firmware version intended for the experiments.",
        "Use the same trigger edge that will be used experimentally.",
        "Keep cable routing and terminations realistic.",
        "Record all settings carefully.",
    ]:
        add_bullet(document, item)
    add_paragraph(document, "Do not attach oscilloscope grounds carelessly. Only probe nodes that are safe for direct oscilloscope connection, and follow the lab grounding rules. If there is any doubt about common ground or floating circuits, use the correct differential or isolated measurement method.")

    document.add_heading("8. Equipment Needed", level=1)
    for item in [
        "M-Duino controller with the intended firmware",
        "Oscilloscope, ideally 4 channels",
        "Suitable probes",
        "BNC cables",
        "BNC splitter",
        "Clean pulse source or function generator",
        "Access to the chosen M-Duino input, the M-Duino DAQ trigger output, DAQ 1 input, DAQ 2 input, and the camera trigger line",
    ]:
        add_bullet(document, item)

    document.add_heading("9. How To Read The Three Diagram Runs", level=1)
    document.add_heading("9.1 Run 1. Controller delay measurement", level=2)
    add_paragraph(document, "Use this setup to measure input to M-Duino DAQ trigger output and input to camera trigger path. It tells you the controller-side delay and jitter.")
    document.add_heading("9.2 Run 2. DAQ splitter timing", level=2)
    add_paragraph(document, "Use this setup to measure M-Duino DAQ output to DAQ 1, M-Duino DAQ output to DAQ 2, and the relative skew between DAQ 1 and DAQ 2.")
    document.add_heading("9.3 Run 3. Camera versus DAQ timing", level=2)
    add_paragraph(document, "Use this setup to measure the relative timing between the camera trigger branch and the DAQ trigger branch.")

    document.add_heading("10. Exactly What To Measure On The Oscilloscope", level=1)
    add_paragraph(document, "Use the same edge definition everywhere. The best practice is to use the real experimental trigger edge and to measure timing at the 50 percent amplitude crossing.")
    for item in [
        "t_input_to_daq = t(DAQ trigger output) - t(input event)",
        "t_input_to_camera = t(camera trigger path) - t(input event)",
        "t_daq1 = t(DAQ 1 trigger input) - t(M-Duino DAQ output)",
        "t_daq2 = t(DAQ 2 trigger input) - t(M-Duino DAQ output)",
        "skew_daq = t(DAQ 1 trigger input) - t(DAQ 2 trigger input)",
        "skew_cam_daq1 = t(camera trigger path) - t(DAQ 1 trigger input)",
        "skew_cam_daq2 = t(camera trigger path) - t(DAQ 2 trigger input)",
    ]:
        add_bullet(document, item)

    document.add_heading("11. Recommended Scope Configuration", level=1)
    for item in [
        "Trigger on the input event for controller-delay measurements.",
        "Trigger on the M-Duino DAQ output for branch comparison measurements.",
        "Use appropriate vertical scale for the logic level.",
        "Use enough time resolution to see the edges cleanly.",
        "Use repeated acquisitions or statistics mode for jitter measurements.",
        "Start with single acquisitions to confirm polarity and signal order.",
    ]:
        add_bullet(document, item)

    document.add_heading("12. Step-By-Step Measurement Procedure", level=1)
    for item in [
        "Confirm that the pulse source produces a clean and repeatable edge.",
        "Confirm that the pulse reaches the intended M-Duino input.",
        "Confirm that the M-Duino produces the expected trigger outputs.",
        "Confirm that the DAQ trigger line and camera trigger line are correctly identified.",
        "Measure controller delay first, then splitter timing, then camera-versus-DAQ timing.",
        "Repeat each measurement over many shots and record statistics.",
    ]:
        add_number(document, item)

    document.add_heading("13. How To Decide If The Measurement Is Valid", level=1)
    for item in [
        "Signals are clean and repeatable.",
        "The same edge definition is used everywhere.",
        "The same reference channel is kept across repeated runs.",
        "The same firmware and mode are used throughout.",
        "The probed points are the real experiment nodes.",
    ]:
        add_bullet(document, item)

    document.add_heading("14. Common Mistakes To Avoid", level=1)
    for item in [
        "Measuring the wrong edge.",
        "Measuring before the real output stage instead of at the actual output node.",
        "Measuring only one DAQ branch and assuming the second one is identical.",
        "Forgetting that the camera is on a different trigger path.",
        "Changing firmware during the test campaign.",
        "Using arbitrary cursor points instead of a consistent crossing definition.",
    ]:
        add_bullet(document, item)

    document.add_heading("15. Results Table", level=1)
    results = document.add_table(rows=8, cols=8)
    headers = ["Path", "Mean delay", "Min", "Max", "Std dev", "Pk-Pk jitter", "Repetitions", "Notes"]
    paths = [
        "Input -> M-Duino DAQ output",
        "Input -> M-Duino camera trigger path",
        "M-Duino DAQ output -> DAQ 1 input",
        "M-Duino DAQ output -> DAQ 2 input",
        "DAQ 1 input -> DAQ 2 input",
        "Camera trigger path -> DAQ 1 input",
        "Camera trigger path -> DAQ 2 input",
    ]
    for c, text in enumerate(headers):
        set_cell_text(results.cell(0, c), text, bold=True)
    for r, path in enumerate(paths, start=1):
        set_cell_text(results.cell(r, 0), path)
    style_table(results)

    document.add_heading("16. Timing Budget Interpretation", level=1)
    add_paragraph(document, "For this application, the M-Duino is mainly a coordination layer, not the final scientific time base.")
    expr = document.add_paragraph()
    run = expr.add_run("Total synchronisation uncertainty ~= controller jitter + branch skew + device trigger-response jitter")
    run.bold = True
    enforce_run_font(run)

    document.add_heading("17. Test Record Template", level=1)
    template = document.add_table(rows=12, cols=2)
    fields = [
        ("Date", ""),
        ("Operator", ""),
        ("Firmware version", ""),
        ("M-Duino mode", ""),
        ("Input channel used", ""),
        ("DAQ trigger output used", ""),
        ("Camera trigger path", ""),
        ("Oscilloscope model", ""),
        ("Probe type", ""),
        ("Trigger edge", ""),
        ("Sample rate", ""),
        ("Repetitions", ""),
    ]
    for r, (label, value) in enumerate(fields):
        set_cell_text(template.cell(r, 0), label, bold=True)
        set_cell_text(template.cell(r, 1), value)
    style_table(template)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            enforce_run_font(run)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        enforce_run_font(run)

    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
