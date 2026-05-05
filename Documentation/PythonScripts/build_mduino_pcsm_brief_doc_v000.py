from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "M-Duino-PCSM_Project_Brief_v000.docx"

ACCENT = RGBColor(27, 76, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "EEF4F8"


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def enforce_run_font(run, font_name="Arial"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def enforce_arial_everywhere(document):
    def process_paragraphs(paragraphs):
        for paragraph in paragraphs:
            for run in paragraph.runs:
                enforce_run_font(run)

    process_paragraphs(document.paragraphs)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs)

    for section in document.sections:
        process_paragraphs(section.header.paragraphs)
        process_paragraphs(section.footer.paragraphs)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    enforce_run_font(run)
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)


def create_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = ACCENT

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(11.5)
    subtitle.font.color.rgb = MUTED

    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_p.add_run("M-Duino-PCSM Project Brief")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    p_pr = header_p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFCAD4")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("Page ")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    set_page_number(footer_p)


def add_paragraph(document, text):
    p = document.add_paragraph()
    run = p.add_run(text)
    enforce_run_font(run)
    return p


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def add_number(document, text):
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def build_document():
    document = Document()
    create_styles(document)
    configure_page(document)

    title = document.add_paragraph(style="Title")
    run = title.add_run("M-Duino-PCSM")
    enforce_run_font(run)

    subtitle = document.add_paragraph(style="Subtitle")
    run = subtitle.add_run("Project brief for the Parameter Control and Status Monitoring application.")
    enforce_run_font(run)

    meta = document.add_table(rows=4, cols=2)
    entries = [
        ("Document date", str(date.today())),
        ("Project title", "M-Duino-PCSM"),
        ("Subtitle", "Parameter Control and Status Monitoring"),
        ("Purpose", "Capture the technical, architectural, and styling decisions required before implementation."),
    ]
    for r, (label, value) in enumerate(entries):
        set_cell_text(meta.cell(r, 0), label, bold=True)
        set_cell_text(meta.cell(r, 1), value)
    style_table(meta)

    document.add_heading("1. Purpose", level=1)
    add_paragraph(document, "Create a desktop or web-style control interface for the M-Duino test controller that allows the operator to change approved parameters between runs, monitor controller status, and view state transitions while preserving the institute visual identity used in the DIME Toolbox.")

    document.add_heading("2. Agreed System Architecture", level=1)
    add_paragraph(document, "The M-Duino remains the real controller. The Python application is only a supervisory layer and must not directly run the real-time firing sequence.")
    add_paragraph(document, "M-Duino responsibilities:")
    for item in [
        "arming logic",
        "firing sequence logic",
        "hot-wire timing",
        "spark timing",
        "DAQ timing",
        "fail, abort, and fired lockout behavior",
        "parameter validation and safety limits",
    ]:
        add_bullet(document, item)
    add_paragraph(document, "PC application responsibilities:")
    for item in [
        "parameter editing",
        "parameter readback",
        "live status display",
        "event and transition log display",
        "serial communication with the M-Duino",
    ]:
        add_bullet(document, item)

    document.add_heading("3. Hardware and Firmware Baseline", level=1)
    add_bullet(document, "Target hardware: Industrial Shields M-Duino 19R+")
    add_bullet(document, "Original reference code: M-DuinoScripts/M-Duino_Original/M-Duino_Original.ino")
    add_bullet(document, "Current cleaned controller code baseline: M-DuinoScripts/M_Duino_v005/M_Duino_v005.ino")
    add_paragraph(document, "The cleaned controller code should be treated as the firmware baseline for app integration because it already includes explicit timing variable names with units, commented sequence logic, named controller states, clearer output/flag separation, debug-oriented state transition logging, spark-test toggle control, and a 30 s spark-test safety timeout.")

    document.add_heading("4. Connection and Communication Assumptions", level=1)
    add_bullet(document, "The app connects to the M-Duino controller itself, not directly to the hot-wire power supply, spark hardware, or DAQ hardware.")
    add_bullet(document, "Primary connection method: USB serial")
    add_bullet(document, "The M-Duino remains powered from its proper external supply.")
    add_bullet(document, "The USB link is used for parameter exchange, status monitoring, and event logging.")

    document.add_heading("5. Critical Interface Requirement", level=1)
    add_paragraph(document, "The app and firmware can only work together if they share a stable external interface. The app does not edit compiled binary values directly. Instead, the firmware must expose a serial command vocabulary that maps to internal variables.")
    add_bullet(document, "Firmware variable example: hotWireBurn_us")
    add_bullet(document, "Protocol name example: HOTWIRE_US")
    add_bullet(document, "Operator-facing label example: Hot-wire burn time")

    document.add_heading("6. Initial Parameter Mapping", level=1)
    table = document.add_table(rows=1, cols=5)
    for idx, header in enumerate(["Firmware variable", "Protocol name", "Operator label", "UI unit", "Notes"]):
        set_cell_text(table.cell(0, idx), header, bold=True)
    rows = [
        ("hotWireBurn_us", "HOTWIRE_US", "Hot-wire burn time", "s", "Main tuning parameter"),
        ("sparkDwell_us", "SPARK_US", "Spark dwell", "us", "Spark pulse duration"),
        ("daqPulse_us", "DAQ_US", "DAQ pulse width", "us", "Must not exceed spark dwell"),
        ("sparkTestInterval_us", "SPARKTEST_US", "Spark-test interval", "ms", "Used in spark-test mode"),
        ("useHotWireStep", "USE_HOTWIRE", "Use hot-wire step", "bool", "Enables or skips melting stage"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    style_table(table)

    document.add_heading("7. Initial Status Contract", level=1)
    add_paragraph(document, "At minimum, the firmware should expose enough information for the app to display:")
    for item in [
        "current controller state",
        "current operating mode",
        "current Arm input state",
        "current Trigger input state",
        "last transition reason",
        "whether the loaded parameter set was accepted",
    ]:
        add_bullet(document, item)
    add_paragraph(document, "Recommended additional status fields:")
    for item in [
        "Spark output active",
        "DAQ output active",
        "hot-wire step enabled",
        "startup configuration summary",
        "firmware version identifier",
    ]:
        add_bullet(document, item)

    document.add_heading("8. Serial Protocol Direction", level=1)
    add_paragraph(document, "The planned communication approach is a simple text-based serial protocol.")
    add_paragraph(document, "Example commands:")
    for cmd in [
        "GET ALL",
        "STATUS?",
        "SET HOTWIRE_US 15000000",
        "SET SPARK_US 5000",
        "SET DAQ_US 600",
        "SET SPARKTEST_US 500000",
        "SET USE_HOTWIRE 1",
    ]:
        add_bullet(document, cmd)
    add_paragraph(document, "Example responses:")
    for rsp in [
        "OK",
        "ERROR NOT_IDLE",
        "ERROR OUT_OF_RANGE",
        "VALUE HOTWIRE_US 15000000",
        "STATE IDLE",
        "MODE HYDROGEN_TEST",
        "ARM 0",
        "TRIGGER 0",
        "REASON operator reset after fired state",
    ]:
        add_bullet(document, rsp)
    add_paragraph(document, "Parameter updates should only be accepted when the controller is in a safe state such as IDLE.")

    document.add_heading("9. App and Firmware Behavior Requirements", level=1)
    add_paragraph(document, "The application should:")
    for item in [
        "detect available serial ports",
        "let the operator choose the correct M-Duino port",
        "connect and disconnect cleanly",
        "request current values on connect",
        "show the latest controller state continuously",
        "show whether a parameter update was accepted or rejected",
        "show the returned reason or error message",
        "keep a readable event log",
        "avoid allowing unsafe parameter edits in the UI",
    ]:
        add_bullet(document, item)
    add_paragraph(document, "The firmware should:")
    for item in [
        "expose the parameter interface over serial",
        "expose live status over serial",
        "expose readable reasons for transitions or rejections",
        "enforce safe-state restrictions for parameter changes",
        "validate all incoming values",
        "keep the firing logic local to the M-Duino",
        "provide a stable naming convention for all exposed values",
    ]:
        add_bullet(document, item)

    document.add_heading("10. Planned Python App Stack", level=1)
    add_bullet(document, "Python")
    add_bullet(document, "Panel")
    add_bullet(document, "pyserial")
    add_paragraph(document, "This stack was chosen because it is open source, flexible, and well suited for a parameter-and-status dashboard.")

    document.add_heading("11. Visual Design Direction", level=1)
    add_paragraph(document, "The app should follow the institute style already used in the DIME Toolbox.")
    for item in [
        "Institute logo included in the header",
        "Inter as the main UI font",
        "JetBrains Mono for technical or console-style text where useful",
        "warm off-white page background",
        "warm light gray surfaces/cards",
        "charcoal text",
        "charcoal top bar",
        "restrained crimson accent",
        "minimal shadows",
        "crisp borders",
        "professional research-tool feel",
    ]:
        add_bullet(document, item)

    document.add_heading("12. Design Tokens Extracted from DIME Toolbox", level=1)
    tokens = document.add_table(rows=1, cols=2)
    for idx, header in enumerate(["Element", "Value"]):
        set_cell_text(tokens.cell(0, idx), header, bold=True)
    token_rows = [
        ("Page background", "#F7F7F5"),
        ("Surface/card", "#EEEEEA"),
        ("Main text", "#333333"),
        ("Muted text", "#888888"),
        ("Top bar", "#2E2E2E"),
        ("Secondary chrome", "#3A3A3A"),
        ("Accent crimson", "#CC2222"),
        ("Accent light", "#F9E5E5"),
        ("Border tone", "#D1D1CD"),
    ]
    for row in token_rows:
        cells = tokens.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    style_table(tokens)

    document.add_heading("13. Proposed App Sections", level=1)
    for item in [
        "Connection panel",
        "Editable parameter panel",
        "Live status panel",
        "Event / transition log",
    ]:
        add_number(document, item)

    document.add_heading("14. First-Version Success Criteria", level=1)
    for item in [
        "Connect to the M-Duino over USB serial",
        "Read the current parameter set",
        "Display the current controller state and core input status",
        "Send a new parameter value while the controller is in IDLE",
        "Receive and display acceptance or rejection from the firmware",
        "Show a readable event or transition log",
        "Preserve the institute visual identity from DIME Toolbox",
    ]:
        add_number(document, item)

    document.add_heading("15. Next Recommended Steps", level=1)
    for item in [
        "Create the GitHub repository for M-Duino-PCSM",
        "Move or copy this brief into that repository",
        "Write the formal serial interface specification",
        "Update the M-Duino firmware to support safe runtime parameter updates",
        "Create the first Python Panel UI scaffold",
        "Add the institute logo and DIME-style theme tokens to the app",
    ]:
        add_number(document, item)

    enforce_arial_everywhere(document)
    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
