from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "M-Duino_Trigger_Timing_Characterisation_Guide_v001.docx"

ACCENT = RGBColor(27, 76, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "EEF4F8"


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def enforce_run_font(run, font_name="Arial"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    enforce_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)


def create_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = ACCENT

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(11.5)
    subtitle.font.color.rgb = MUTED

    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_p.add_run("M-Duino Trigger Timing Characterisation Guide")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    p_pr = header_p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFCAD4")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("Page ")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    set_page_number(footer_p)


def add_paragraph(document, text, bold=False):
    p = document.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    enforce_run_font(run)
    return p


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def add_number(document, text):
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def add_wiring_table(document, title_text, rows):
    document.add_heading(title_text, level=2)
    table = document.add_table(rows=len(rows) + 1, cols=3)
    headers = ["Connection point", "Connect to", "Why it is measured"]
    for c, text in enumerate(headers):
        set_cell_text(table.cell(0, c), text, bold=True)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            set_cell_text(table.cell(r, c), text)
    style_table(table)


def build_document():
    document = Document()
    create_styles(document)
    configure_page(document)

    title = document.add_paragraph(style="Title")
    run = title.add_run("M-Duino Trigger Timing Characterisation Guide")
    enforce_run_font(run)

    subtitle = document.add_paragraph(style="Subtitle")
    run = subtitle.add_run("Detailed oscilloscope procedure for delay, jitter, and synchronisation checks.")
    enforce_run_font(run)

    meta = document.add_table(rows=4, cols=2)
    entries = [
        ("Document date", str(date.today())),
        ("Application", "M-Duino trigger coordination"),
        ("Primary use", "Deflagration trigger-path characterisation"),
        ("Focus", "Delay, jitter, and relative skew validation"),
    ]
    for r, (label, value) in enumerate(entries):
        set_cell_text(meta.cell(r, 0), label, bold=True)
        set_cell_text(meta.cell(r, 1), value)
    style_table(meta)

    document.add_heading("1. Purpose", level=1)
    add_paragraph(document, "This guide explains how to measure the timing behaviour of the M-Duino when it is used as a trigger coordinator for a hydrogen deflagration setup. The objective is to quantify controller delay, shot-to-shot jitter, and the relative synchronisation offsets between the DAQ trigger branch and the camera trigger branch.")

    document.add_heading("2. What This Characterisation Is Trying To Prove", level=1)
    add_paragraph(document, "The M-Duino is not the final scientific timing reference in this setup. The final scientific time base is provided by the high-speed datalogger. The main question is therefore whether the M-Duino distributes the trigger events with low enough uncertainty for the experiment.")
    for item in [
        "Low and stable trigger delay",
        "Low trigger jitter",
        "Low relative skew between systems",
    ]:
        add_bullet(document, item)

    document.add_heading("3. Assumed Trigger Architecture", level=1)
    for item in [
        "A clean input event is applied to the M-Duino.",
        "The M-Duino produces a DAQ trigger output.",
        "That DAQ trigger is split through a BNC splitter to DAQ 1 and DAQ 2.",
        "The camera is triggered through its own separate trigger path.",
        "Final scientific timestamping is performed by the high-speed datalogger.",
    ]:
        add_number(document, item)
    add_paragraph(document, "There are therefore two timing branches to characterise: the split DAQ trigger branch and the separate camera trigger branch.")

    document.add_heading("4. Why Each Measurement Is Needed", level=1)
    add_paragraph(document, "Each timing measurement answers a different question.")
    for item in [
        "Input to M-Duino DAQ trigger output: controller reaction delay and jitter.",
        "Input to camera trigger path: controller-to-camera path timing and jitter.",
        "M-Duino DAQ output to DAQ 1 and DAQ 2 inputs: splitter-path delay and branch-to-branch comparison.",
        "Camera trigger path relative to DAQ trigger path: inter-system synchronisation offset and stability.",
    ]:
        add_bullet(document, item)

    document.add_heading("5. Do You Need Special Benchmark Code?", level=1)
    add_paragraph(document, "Not necessarily for the first pass. A real-system timing study with the actual firmware and actual trigger paths is usually the most important measurement, because it answers whether the experiment-ready system is good enough.")
    add_paragraph(document, "Special benchmark firmware of the type shown in the Industrial Shields article is optional and mainly useful if the real-system timing looks worse than expected or if the delay needs to be decomposed into hardware and firmware contributions.")

    document.add_heading("6. Safety and Measurement Validity", level=1)
    for item in [
        "Perform the test without ignition and without hazardous experiment execution.",
        "Use a safe and repeatable pulse source.",
        "Use the same firmware version intended for the experiments.",
        "Use the same trigger edge that will be used experimentally.",
        "Keep cable routing and terminations realistic.",
        "Record all settings carefully.",
    ]:
        add_bullet(document, item)
    add_paragraph(document, "Do not attach oscilloscope grounds carelessly. Only probe nodes that are safe for direct oscilloscope connection, and follow the lab grounding rules. If there is any doubt about common ground or floating circuits, use the correct differential or isolated measurement method.")

    document.add_heading("7. Equipment Needed", level=1)
    for item in [
        "M-Duino controller with the intended firmware",
        "Oscilloscope, ideally with four channels",
        "Probes suitable for the signal levels being measured",
        "BNC cables and BNC splitter",
        "Clean pulse source or function generator",
        "Access to the chosen M-Duino input, M-Duino DAQ trigger output, DAQ 1 input, DAQ 2 input, and camera trigger line",
    ]:
        add_bullet(document, item)

    document.add_heading("8. Recommended Wiring Concept", level=1)
    add_wiring_table(
        document,
        "8.1 Controller delay measurement",
        [
            ["Pulse source output", "M-Duino input under test", "Creates the test event"],
            ["Pulse source output", "Oscilloscope CH1", "Reference timing event"],
            ["M-Duino DAQ trigger output", "Oscilloscope CH2", "Measures controller DAQ output delay"],
            ["M-Duino camera trigger path", "Oscilloscope CH3", "Measures controller camera-path delay"],
        ],
    )
    add_wiring_table(
        document,
        "8.2 DAQ splitter path measurement",
        [
            ["M-Duino DAQ trigger output before splitter", "Oscilloscope CH1", "Common reference for branch timing"],
            ["Splitter output branch 1 / DAQ 1 input", "Oscilloscope CH2", "Measures DAQ 1 arrival time"],
            ["Splitter output branch 2 / DAQ 2 input", "Oscilloscope CH3", "Measures DAQ 2 arrival time"],
        ],
    )
    add_wiring_table(
        document,
        "8.3 Camera versus DAQ timing measurement",
        [
            ["M-Duino DAQ trigger output", "Oscilloscope CH1", "Reference timing event"],
            ["DAQ trigger input", "Oscilloscope CH2", "DAQ branch timing"],
            ["Camera trigger path", "Oscilloscope CH3", "Camera branch timing"],
        ],
    )

    document.add_heading("9. Recommended Scope Channels", level=1)
    add_paragraph(document, "If four channels are available, use separate runs for controller delay, DAQ splitter timing, and camera versus DAQ timing. If not all branches can be measured at once, keep one common reference channel across all runs. The best common reference is usually the M-Duino DAQ trigger output before the splitter.")

    document.add_heading("10. Exactly What To Measure On The Oscilloscope", level=1)
    add_paragraph(document, "All timing measurements should use the same edge definition. The best practice is to use the actual experimental trigger edge and to measure timing at the 50 percent amplitude crossing of that edge.")
    add_paragraph(document, "If the oscilloscope supports automatic delay measurements, use them. If not, place the time cursors at the 50 percent crossing points and use the same method for every signal.")
    for item in [
        "t_input_to_daq = t(DAQ trigger output) - t(input event)",
        "t_input_to_camera = t(camera trigger path) - t(input event)",
        "t_daq1 = t(DAQ 1 trigger input) - t(M-Duino DAQ output)",
        "t_daq2 = t(DAQ 2 trigger input) - t(M-Duino DAQ output)",
        "skew_daq = t(DAQ 1 trigger input) - t(DAQ 2 trigger input)",
        "skew_cam_daq1 = t(camera trigger path) - t(DAQ 1 trigger input)",
        "skew_cam_daq2 = t(camera trigger path) - t(DAQ 2 trigger input)",
    ]:
        add_bullet(document, item)

    document.add_heading("11. Recommended Scope Configuration", level=1)
    for item in [
        "Trigger on the input event for controller-delay measurements.",
        "Trigger on the M-Duino DAQ output for branch comparison measurements.",
        "Use vertical scales appropriate for the logic level.",
        "Use enough time resolution to see the edges cleanly.",
        "Use repeated acquisitions or statistics mode for jitter measurements.",
        "Start with single acquisitions to confirm signal order and polarity.",
    ]:
        add_bullet(document, item)

    document.add_heading("12. Step-By-Step Measurement Procedure", level=1)
    document.add_heading("12.1 Validate the bench setup", level=2)
    for item in [
        "Confirm that the pulse source produces a clean and repeatable edge.",
        "Confirm that the pulse actually reaches the intended M-Duino input.",
        "Confirm that the M-Duino produces the expected trigger outputs.",
        "Confirm that the camera trigger line and DAQ trigger line are both visible and correctly identified.",
    ]:
        add_number(document, item)
    document.add_heading("12.2 Measure controller delay", level=2)
    for item in [
        "Connect the pulse source to the chosen M-Duino input.",
        "Probe the input event on CH1.",
        "Probe the M-Duino DAQ trigger output on CH2.",
        "Probe the camera trigger path on CH3.",
        "Trigger the scope on the input event.",
        "Measure input to DAQ trigger output and input to camera trigger path.",
        "Repeat for at least 100 shots, preferably 500 to 1000 if the scope statistics mode supports it.",
    ]:
        add_number(document, item)
    document.add_heading("12.3 Measure DAQ splitter timing", level=2)
    for item in [
        "Probe the M-Duino DAQ trigger output before the splitter on CH1.",
        "Probe DAQ 1 trigger input on CH2.",
        "Probe DAQ 2 trigger input on CH3.",
        "Trigger the scope on the M-Duino DAQ output.",
        "Measure DAQ output to DAQ 1, DAQ output to DAQ 2, and DAQ 1 to DAQ 2 relative skew.",
        "Repeat for many shots and record statistics.",
    ]:
        add_number(document, item)
    document.add_heading("12.4 Measure camera versus DAQ timing", level=2)
    for item in [
        "Keep the M-Duino DAQ trigger output as the common reference.",
        "Probe the camera trigger path.",
        "Probe at least one DAQ trigger input.",
        "Measure the relative delay between the camera path and the DAQ path.",
        "Repeat for many shots and record statistics.",
    ]:
        add_number(document, item)

    document.add_heading("13. How To Decide If The Measurement Is Valid", level=1)
    for item in [
        "Signals are clean and repeatable.",
        "The same edge definition is used everywhere.",
        "The reference channel is kept consistent across repeated runs.",
        "The same firmware and mode are used throughout the test.",
        "The probing points are the real experiment nodes, not merely convenient intermediate points.",
    ]:
        add_bullet(document, item)

    document.add_heading("14. Common Mistakes To Avoid", level=1)
    for item in [
        "Measuring the wrong trigger edge.",
        "Measuring before the real output stage instead of at the actual output node.",
        "Measuring only one DAQ branch and assuming the second one is identical.",
        "Forgetting that the camera is on a different trigger path.",
        "Changing firmware or controller configuration during the test campaign.",
        "Using arbitrary cursor positions instead of a consistent edge crossing definition.",
        "Failing to record which physical node was probed.",
    ]:
        add_bullet(document, item)

    document.add_heading("15. Statistics To Record", level=1)
    for item in [
        "Mean delay",
        "Minimum delay",
        "Maximum delay",
        "Standard deviation",
        "Peak-to-peak jitter",
        "Number of repetitions",
    ]:
        add_bullet(document, item)

    document.add_heading("16. Results Table", level=1)
    results = document.add_table(rows=8, cols=8)
    headers = [
        "Path",
        "Mean delay",
        "Min",
        "Max",
        "Std dev",
        "Pk-Pk jitter",
        "Repetitions",
        "Notes",
    ]
    paths = [
        "Input -> M-Duino DAQ output",
        "Input -> M-Duino camera trigger path",
        "M-Duino DAQ output -> DAQ 1 input",
        "M-Duino DAQ output -> DAQ 2 input",
        "DAQ 1 input -> DAQ 2 input",
        "Camera trigger path -> DAQ 1 input",
        "Camera trigger path -> DAQ 2 input",
    ]
    for c, text in enumerate(headers):
        set_cell_text(results.cell(0, c), text, bold=True)
    for r, path in enumerate(paths, start=1):
        set_cell_text(results.cell(r, 0), path)
    style_table(results)

    document.add_heading("17. Timing Budget Interpretation", level=1)
    add_paragraph(document, "For this application, the M-Duino acts mainly as an event coordinator rather than the final scientific timing reference. The timing budget should therefore focus on delay stability, jitter, and relative branch skew.")
    expr = document.add_paragraph()
    run = expr.add_run("Total synchronisation uncertainty ~= controller jitter + branch skew + device trigger-response jitter")
    run.bold = True
    enforce_run_font(run)
    add_paragraph(document, "For deflagration work, a stable fixed offset is often much less problematic than a variable shot-to-shot offset.")

    document.add_heading("18. Suggested Acceptance Logic", level=1)
    for item in [
        "DAQ 1 and DAQ 2 receive the split trigger with negligible relative skew for the intended experiment.",
        "The camera path has a stable offset relative to the DAQ path.",
        "M-Duino-trigger jitter is small compared with the timing precision required by the experiment.",
        "The total uncertainty is small relative to the characteristic deflagration timescales of interest.",
    ]:
        add_bullet(document, item)

    document.add_heading("19. Example Report Wording", level=1)
    add_paragraph(document, "The trigger timing of the M-Duino coordination layer was characterised with an oscilloscope prior to experiments. The delay from the selected M-Duino input event to the DAQ trigger output and to the camera trigger path was measured over repeated trials, together with the relative skew between the two DAQ trigger inputs and the camera trigger path. Because the final scientific timestamping was performed by the high-speed datalogger, the primary acceptance criterion was low and stable inter-system trigger uncertainty rather than absolute controller latency. For the present deflagration experiments, the measured trigger offsets were sufficiently stable that inter-device clock differences were considered negligible over the acquisition window.")

    document.add_heading("20. Test Record Template", level=1)
    template = document.add_table(rows=12, cols=2)
    fields = [
        ("Date", ""),
        ("Operator", ""),
        ("Firmware version", ""),
        ("M-Duino mode", ""),
        ("Input channel used", ""),
        ("DAQ trigger output used", ""),
        ("Camera trigger path", ""),
        ("Oscilloscope model", ""),
        ("Probe type", ""),
        ("Trigger edge", ""),
        ("Sample rate", ""),
        ("Repetitions", ""),
    ]
    for r, (label, value) in enumerate(fields):
        set_cell_text(template.cell(r, 0), label, bold=True)
        set_cell_text(template.cell(r, 1), value)
    style_table(template)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            enforce_run_font(run)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        enforce_run_font(run)

    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
