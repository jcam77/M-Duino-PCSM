from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "M-Duino_Trigger_Timing_Characterisation_Guide_v000.docx"

ACCENT = RGBColor(27, 76, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "EEF4F8"


def set_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def enforce_run_font(run, font_name="Arial"):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    enforce_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)


def create_styles(document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = ACCENT

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(11.5)
    subtitle.font.color.rgb = MUTED

    for name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header_p.add_run("M-Duino Trigger Timing Characterisation Guide")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    p_pr = header_p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFCAD4")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_p.add_run("Page ")
    enforce_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    set_page_number(footer_p)


def add_paragraph(document, text):
    p = document.add_paragraph()
    run = p.add_run(text)
    enforce_run_font(run)
    return p


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def add_number(document, text):
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    enforce_run_font(run)


def build_document():
    document = Document()
    create_styles(document)
    configure_page(document)

    title = document.add_paragraph(style="Title")
    run = title.add_run("M-Duino Trigger Timing Characterisation Guide")
    enforce_run_font(run)

    subtitle = document.add_paragraph(style="Subtitle")
    run = subtitle.add_run("Oscilloscope procedure for delay, jitter, and synchronisation checks.")
    enforce_run_font(run)

    meta = document.add_table(rows=4, cols=2)
    entries = [
        ("Document date", str(date.today())),
        ("Application", "M-Duino trigger coordination"),
        ("Primary use", "Deflagration trigger-path characterisation"),
        ("Main outputs", "Delay, jitter, and relative skew measurements"),
    ]
    for r, (label, value) in enumerate(entries):
        set_cell_text(meta.cell(r, 0), label, bold=True)
        set_cell_text(meta.cell(r, 1), value)
    style_table(meta)

    document.add_heading("1. Purpose", level=1)
    add_paragraph(document, "This guide explains how to characterise the trigger timing performance of the M-Duino when it is used as an event sequencer for hydrogen deflagration experiments. The objective is to measure controller delay, shot-to-shot jitter, and the relative synchronisation offsets between the DAQ and camera trigger paths.")

    document.add_heading("2. Assumed Trigger Architecture", level=1)
    add_number(document, "An input event arrives at the M-Duino.")
    add_number(document, "The M-Duino generates a DAQ trigger output.")
    add_number(document, "That DAQ trigger is split through a BNC splitter to DAQ 1 and DAQ 2.")
    add_number(document, "The camera is triggered through its own separate trigger path.")
    add_number(document, "Final scientific timestamping is performed by the high-speed datalogger.")

    document.add_heading("3. What Must Be Measured", level=1)
    add_bullet(document, "Input event to M-Duino DAQ trigger output delay")
    add_bullet(document, "Input event to camera trigger path delay")
    add_bullet(document, "DAQ branch delays after the splitter")
    add_bullet(document, "Relative skew between DAQ 1 and DAQ 2")
    add_bullet(document, "Relative skew between the camera trigger path and the DAQ trigger path")
    add_bullet(document, "Shot-to-shot jitter for all relevant timing paths")

    document.add_heading("4. Safety and Test Conditions", level=1)
    add_bullet(document, "Perform the timing check without ignition or hazardous experiment execution.")
    add_bullet(document, "Use a safe and repeatable input pulse source.")
    add_bullet(document, "Keep the firmware version fixed during the measurements.")
    add_bullet(document, "Use the same trigger edge as the real experiment.")
    add_bullet(document, "Record cable lengths, scope settings, and all signal names.")

    document.add_heading("5. Recommended Equipment", level=1)
    for item in [
        "M-Duino controller with the intended test firmware",
        "Oscilloscope, preferably with four channels",
        "BNC cables and splitter",
        "Safe pulse source or function generator",
        "Access to the M-Duino input, DAQ trigger output, DAQ trigger inputs, and camera trigger line",
    ]:
        add_bullet(document, item)

    document.add_heading("6. Suggested Channel Assignment", level=1)
    table = document.add_table(rows=5, cols=3)
    headers = ["Scope channel", "Signal", "Purpose"]
    rows = [
        ["CH1", "M-Duino input event", "Timing reference"],
        ["CH2", "M-Duino DAQ trigger output", "Controller output delay"],
        ["CH3", "DAQ 1 or DAQ 2 trigger input", "Distribution-path timing"],
        ["CH4", "Camera trigger signal or camera trigger input", "Camera-to-DAQ relative timing"],
    ]
    for c, text in enumerate(headers):
        set_cell_text(table.cell(0, c), text, bold=True)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            set_cell_text(table.cell(r, c), text)
    style_table(table)
    add_paragraph(document, "Repeat the DAQ measurement with DAQ 1 and DAQ 2 in separate runs if all signals cannot be observed at once.")

    document.add_heading("7. Test Procedure", level=1)
    document.add_heading("7.1 Bench sanity check", level=2)
    for item in [
        "Confirm all channels are connected correctly.",
        "Verify the pulse source is safe and repeatable.",
        "Confirm the M-Duino responds as expected.",
        "Verify that both the DAQ trigger path and the camera trigger path are visible on the oscilloscope.",
    ]:
        add_number(document, item)

    document.add_heading("7.2 M-Duino delay measurement", level=2)
    for item in [
        "Connect the pulse source to the M-Duino input under test.",
        "Trigger the scope on the selected input edge.",
        "Measure the delay from the input event to the M-Duino DAQ trigger output.",
        "Measure the delay from the input event to the camera trigger path.",
        "Repeat for at least 100 events, preferably 500 to 1000 if statistics mode is available.",
    ]:
        add_number(document, item)

    document.add_heading("7.3 DAQ splitter path measurement", level=2)
    for item in [
        "Observe the M-Duino DAQ trigger output before the splitter.",
        "Observe the DAQ trigger inputs after the splitter.",
        "Measure the delay to each DAQ branch.",
        "Measure the relative skew between DAQ 1 and DAQ 2.",
        "Repeat for many events and record statistics.",
    ]:
        add_number(document, item)

    document.add_heading("7.4 Camera-to-DAQ synchronisation measurement", level=2)
    for item in [
        "Observe the camera trigger line together with the DAQ trigger path.",
        "Measure the relative delay between the camera trigger path and the DAQ trigger path.",
        "Repeat for many events and record statistics.",
        "If required, repeat with both DAQ channels in separate runs.",
    ]:
        add_number(document, item)

    document.add_heading("8. Statistics To Record", level=1)
    for item in [
        "Mean delay",
        "Minimum delay",
        "Maximum delay",
        "Standard deviation",
        "Peak-to-peak jitter",
        "Number of repetitions",
    ]:
        add_bullet(document, item)

    document.add_heading("9. Results Table", level=1)
    results = document.add_table(rows=8, cols=8)
    headers = [
        "Path",
        "Mean delay",
        "Min",
        "Max",
        "Std dev",
        "Pk-Pk jitter",
        "Repetitions",
        "Notes",
    ]
    paths = [
        "Input -> M-Duino DAQ output",
        "Input -> M-Duino camera trigger path",
        "M-Duino DAQ output -> DAQ 1 input",
        "M-Duino DAQ output -> DAQ 2 input",
        "DAQ 1 input -> DAQ 2 input",
        "Camera trigger path -> DAQ 1 input",
        "Camera trigger path -> DAQ 2 input",
    ]
    for c, text in enumerate(headers):
        set_cell_text(results.cell(0, c), text, bold=True)
    for r, path in enumerate(paths, start=1):
        set_cell_text(results.cell(r, 0), path)
    style_table(results)

    document.add_heading("10. Timing Budget Interpretation", level=1)
    add_paragraph(document, "For this application, the M-Duino acts mainly as an event coordinator rather than the final scientific timing reference. The timing budget should therefore focus on delay stability, jitter, and relative skew between systems.")
    add_paragraph(document, "A practical expression is:")
    expr = document.add_paragraph()
    run = expr.add_run("Total synchronisation uncertainty ~= controller jitter + branch skew + device trigger-response jitter")
    run.bold = True
    enforce_run_font(run)
    add_paragraph(document, "A fixed delay is often less problematic than a variable delay because a stable fixed offset can often be corrected during analysis.")

    document.add_heading("11. Example Report Wording", level=1)
    add_paragraph(document, "The trigger timing of the M-Duino coordination layer was characterised with an oscilloscope prior to experiments. The delay from the selected M-Duino input event to the DAQ trigger output and camera trigger path was measured over repeated trials, together with the relative skew between the two DAQ trigger inputs and the camera trigger path. Because the final scientific timestamping was performed by the high-speed datalogger, the main acceptance criterion was low and stable inter-system trigger uncertainty rather than absolute controller latency. For the present deflagration experiments, the measured trigger offsets were sufficiently stable that inter-device clock differences were considered negligible over the acquisition window.")

    document.add_heading("12. Test Record Template", level=1)
    template = document.add_table(rows=11, cols=2)
    fields = [
        ("Date", ""),
        ("Operator", ""),
        ("Firmware version", ""),
        ("M-Duino mode", ""),
        ("Input channel used", ""),
        ("DAQ trigger output used", ""),
        ("Camera trigger output/path", ""),
        ("Oscilloscope model", ""),
        ("Trigger edge", ""),
        ("Sample rate", ""),
        ("Repetitions", ""),
    ]
    for r, (label, value) in enumerate(fields):
        set_cell_text(template.cell(r, 0), label, bold=True)
        set_cell_text(template.cell(r, 1), value)
    style_table(template)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            enforce_run_font(run)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        enforce_run_font(run)

    document.save(DOCX_PATH)


if __name__ == "__main__":
    build_document()
